"""Export FAQ lists to CSV, Excel, Word, or PDF."""

from __future__ import annotations

import csv
import io
import logging
import re
from typing import Any
from xml.sax.saxutils import escape as xml_escape

logger = logging.getLogger(__name__)

_HTML_TAG = re.compile(r"<[^>]+>", re.I)
_ANCHOR = re.compile(
    r'<a\s[^>]*\bhref\s*=\s*["\']([^"\']+)["\'][^>]*>([\s\S]*?)</a>',
    re.I,
)


def _plain_segment_preserve_boundaries(s: str) -> str:
    """
    Strip tags and collapse internal whitespace, but keep one leading / trailing space when
    the source had whitespace there (so spaces next to <a> tags are not lost).
    """
    if not s:
        return ""
    t = _HTML_TAG.sub("", s)
    lead = bool(re.match(r"^\s", t))
    trail = bool(re.search(r"\s$", t))
    inner = t.strip()
    if not inner:
        return " " if (lead or trail) else ""
    middle = " ".join(inner.split())
    return ((" " if lead else "") + middle + (" " if trail else ""))


def _join_segments_need_space(prev: str, nxt: str) -> bool:
    """True if a space should be inserted between two consecutive exported fragments."""
    if not prev or not nxt:
        return False
    return not (prev[-1].isspace() or nxt[0].isspace())


def iter_answer_segments(html: str):
    """
    Yield ('plain', text_chunk) or ('link', url, label).
    `html` is FAQ answer text; internal links use <a href="https://...">label</a>.
    """
    if not html:
        return
    pos = 0
    for m in _ANCHOR.finditer(html):
        if m.start() > pos:
            yield ("plain", html[pos : m.start()])
        url = (m.group(1) or "").strip()
        label = _HTML_TAG.sub("", m.group(2) or "").strip()
        yield ("link", url, label)
        pos = m.end()
    if pos < len(html):
        yield ("plain", html[pos:])


def answer_to_plain_export(text: str) -> str:
    """Plain text for CSV / spreadsheet cells (labels only; no pasted URLs)."""
    if not text:
        return ""
    parts: list[str] = []
    for seg in iter_answer_segments(text):
        if seg[0] == "plain":
            p = _plain_segment_preserve_boundaries(seg[1])
            if p:
                parts.append(p)
        else:
            _url, label = seg[1], seg[2]
            if label:
                parts.append(label)
    out: list[str] = []
    for p in parts:
        if not p:
            continue
        if out and _join_segments_need_space(out[-1], p):
            out.append(" ")
        out.append(p)
    return "".join(out).strip()


def answer_to_reportlab_markup(answer_html: str) -> str:
    """Mini-HTML for ReportLab Paragraph: clickable <a href> links, rest escaped."""
    if not answer_html or not str(answer_html).strip():
        return "—"
    chunks: list[str] = []
    prev_tail = ""
    for seg in iter_answer_segments(answer_html):
        if seg[0] == "plain":
            p = _plain_segment_preserve_boundaries(seg[1])
            if not p:
                continue
            if prev_tail and _join_segments_need_space(prev_tail, p):
                chunks.append(" ")
            chunks.append(xml_escape(p))
            prev_tail = p
        else:
            url, label = seg[1], seg[2]
            if not url:
                continue
            disp = label if label else "Link"
            if prev_tail and _join_segments_need_space(prev_tail, disp):
                chunks.append(" ")
            chunks.append(
                '<a href="%s" color="blue">%s</a>'
                % (xml_escape(url, {"'": "&apos;"}), xml_escape(disp))
            )
            prev_tail = disp
    return "".join(chunks) if chunks else "—"


def _docx_add_hyperlink(paragraph, url: str, text: str) -> None:
    """Append an external hyperlink run (Word clickable link)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(Pt(10) * 2)))
    r_pr.append(sz)
    new_run.append(r_pr)

    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _docx_append_answer_paragraph(doc, answer_html: str) -> None:
    """Answer paragraph with preserved hyperlinks (no raw URLs unless model pasted them)."""
    from docx.shared import Pt

    p = doc.add_paragraph()
    empty = True
    prev_tail = ""
    for seg in iter_answer_segments(answer_html):
        if seg[0] == "plain":
            chunk = _plain_segment_preserve_boundaries(seg[1])
            if not chunk:
                continue
            if prev_tail and _join_segments_need_space(prev_tail, chunk):
                gap = p.add_run(" ")
                gap.font.size = Pt(10)
            run = p.add_run(chunk)
            run.font.size = Pt(10)
            prev_tail = chunk
            empty = False
        else:
            url, label = seg[1], seg[2]
            if not url:
                continue
            disp = label if label else "Link"
            if prev_tail and _join_segments_need_space(prev_tail, disp):
                gap = p.add_run(" ")
                gap.font.size = Pt(10)
            _docx_add_hyperlink(p, url, disp)
            prev_tail = disp
            empty = False
    if empty:
        run = p.add_run("—")
        run.font.size = Pt(10)
    for run in p.runs:
        if run.font.size is None:
            run.font.size = Pt(10)


def _answer_html_str(raw: Any) -> str:
    if raw is None:
        return ""
    return raw if isinstance(raw, str) else str(raw)


def _safe_filename(title: str, ext: str) -> str:
    base = re.sub(r"[^\w\s\-]", "", (title or "FAQs").strip())[:80] or "FAQs"
    base = re.sub(r"\s+", "_", base)
    return f"{base}.{ext}"


def export_faqs_bytes(
    fmt: str,
    faqs: list[dict[str, Any]],
    title: str = "FAQs",
) -> tuple[bytes, str, str]:
    """
    Build file bytes for the given format.
    Returns (body_bytes, download_filename, media_type).
    """
    fmt = (fmt or "csv").lower().strip()
    if fmt == "csv":
        return _export_csv(faqs, title)
    if fmt == "xlsx":
        return _export_xlsx(faqs, title)
    if fmt == "docx":
        return _export_docx(faqs, title)
    if fmt == "pdf":
        return _export_pdf(faqs, title)
    raise ValueError(f"Unsupported export format: {fmt}")


def _export_csv(faqs: list[dict[str, Any]], title: str) -> tuple[bytes, str, str]:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Question", "Answer"])
    for f in faqs:
        writer.writerow(
            [
                (f.get("question") or "").replace("\r\n", "\n"),
                answer_to_plain_export(f.get("answer") or "").replace("\r\n", "\n"),
            ]
        )
    data = buf.getvalue().encode("utf-8-sig")
    return data, _safe_filename(title, "csv"), "text/csv; charset=utf-8"


def _export_xlsx(faqs: list[dict[str, Any]], title: str) -> tuple[bytes, str, str]:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "FAQs"
    ws.append(["Question", "Answer"])
    for f in faqs:
        ws.append(
            [
                (f.get("question") or ""),
                answer_to_plain_export(f.get("answer") or ""),
            ]
        )
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 72
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), _safe_filename(title, "xlsx"), (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def _export_docx(faqs: list[dict[str, Any]], title: str) -> tuple[bytes, str, str]:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title or "FAQs", level=0)
    for i, f in enumerate(faqs, start=1):
        q = f.get("question") or ""
        p = doc.add_paragraph()
        run = p.add_run(f"Q{i}. {q}")
        run.bold = True
        run.font.size = Pt(11)
        _docx_append_answer_paragraph(doc, _answer_html_str(f.get("answer")))
        doc.add_paragraph()
    out = io.BytesIO()
    doc.save(out)
    return (
        out.getvalue(),
        _safe_filename(title, "docx"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def _export_pdf(faqs: list[dict[str, Any]], title: str) -> tuple[bytes, str, str]:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=inch * 0.75,
        leftMargin=inch * 0.75,
        topMargin=inch * 0.75,
        bottomMargin=inch * 0.75,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "FAQTitle",
        parent=styles["Heading1"],
        fontSize=16,
        spaceAfter=14,
    )
    q_style = ParagraphStyle(
        "FAQQuestion",
        parent=styles["Normal"],
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#1e40af"),
        spaceBefore=10,
        spaceAfter=4,
        fontName="Helvetica-Bold",
    )
    a_style = ParagraphStyle(
        "FAQAnswer",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
        spaceAfter=8,
        fontName="Helvetica",
    )
    story = []
    story.append(Paragraph(xml_escape(title or "FAQs"), title_style))
    story.append(Spacer(1, 6))
    for i, f in enumerate(faqs, start=1):
        q = xml_escape((f.get("question") or "").strip() or "—")
        a_markup = answer_to_reportlab_markup(_answer_html_str(f.get("answer")))
        story.append(Paragraph(f"Q{i}. {q}", q_style))
        story.append(Paragraph(a_markup, a_style))
    doc.build(story)
    return buffer.getvalue(), _safe_filename(title, "pdf"), "application/pdf"
